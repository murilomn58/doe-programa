"""Tests for table + report export (issue 09).

Three layers:
1. Pure writer unit tests — feed the CSV/XLSX writers a real
   ``AnalysisResultData`` built from the golden 2^2 data and assert the PT-BR
   conventions (BOM, ``;``, ``,``), the ANOVA numbers, and the 3 XLSX sheets.
2. HTTP end-to-end via ``AsyncClient`` — seed the class 2^2 design and POST
   ``/export/table``, ``/export/xlsx`` -> 200 with the right content types.
3. PDF report — render a minimal local ``file://`` page carrying ``#report-ready``
   to a real PDF via Playwright; skipped (not failed) if Chromium is missing.
"""

import io
import tempfile
from pathlib import Path

import openpyxl
import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from db import get_session
from domain.models import Base, Factor, Response, Run
from domain.types import RSMModel
from main import app
from services.analysis_service import FactorInfo, analyze, terms_for_model
from services.export_service import (
    build_report_pdf,
    coefficients_table,
    design_matrix_table,
    to_csv_ptbr,
    to_xlsx,
)

# --- golden 2^2 fixture (same numbers as test_publication_export) ------------

GOLDEN = [
    (1, -1, 3.475),
    (-1, -1, 0.612),
    (1, 1, 9.041),
    (-1, 1, 5.168),
    (-1, -1, 0.645),
    (1, -1, 1.623),
    (1, 1, 6.27),
    (-1, 1, 3.447),
]
NAT = {"cond": (18.0, 19.0), "par": (70.0, 210.0)}


def _to_nat(name: str, coded: float) -> float:
    low, high = NAT[name]
    return low + (coded + 1.0) / 2.0 * (high - low)


def _golden_result():
    factors = [
        FactorInfo("cond", "quantitative"),
        FactorInfo("par", "quantitative"),
    ]
    coded = [{"cond": c, "par": p} for c, p, _ in GOLDEN]
    natural = [
        {"cond": _to_nat("cond", c), "par": _to_nat("par", p)} for c, p, _ in GOLDEN
    ]
    y = [v for _, _, v in GOLDEN]
    terms = terms_for_model(RSMModel.LINEAR, factors)
    return analyze(
        factors=factors,
        coded=coded,
        natural=natural,
        y=y,
        terms=terms,
        model_label="linear",
        response_name="DP_50000",
    )


def _orm_fixture():
    """Minimal ORM objects for the design-matrix / xlsx writers (no DB)."""
    factors = [
        Factor(name="cond", factor_type="quantitative", unit=None, position=0),
        Factor(name="par", factor_type="quantitative", unit=None, position=1),
    ]
    responses = [Response(name="DP_50000", unit="mm", position=0)]
    runs: list[Run] = []
    for i, (c, p, y) in enumerate(GOLDEN):
        runs.append(
            Run(
                run_order=i + 1,
                std_order=i + 1,
                point_type=1,
                block=1,
                replicate=1,
                cells={
                    "cond": {"coded": c, "natural": _to_nat("cond", c)},
                    "par": {"coded": p, "natural": _to_nat("par", p)},
                },
                response_values={"DP_50000": y},
            )
        )
    return factors, responses, runs


# --- CSV PT-BR ---------------------------------------------------------------


def test_csv_starts_with_bom_and_uses_ptbr_separators() -> None:
    data = to_csv_ptbr([[1, 2.5, "x"]], ["a", "b", "c"])
    assert data[:3] == b"\xef\xbb\xbf"  # utf-8-sig BOM
    text = data.decode("utf-8-sig")
    assert ";" in text
    assert "2,5" in text  # decimal comma, not point
    assert "." not in text.split("\r\n")[1]  # no decimal points in the data row


def test_csv_quotes_fields_containing_separator() -> None:
    data = to_csv_ptbr([["has;semi", "plain"]], ["h1", "h2"])
    text = data.decode("utf-8-sig")
    assert '"has;semi"' in text


def test_csv_anova_rows_match_analysis() -> None:
    from services.export_service import anova_table

    result = _golden_result()
    headers, rows = anova_table(result)
    data = to_csv_ptbr(rows, headers)
    text = data.decode("utf-8-sig")
    lines = text.strip().split("\r\n")
    assert lines[0] == "Fonte;SQ;GL;QM;F;p-valor"
    # The ANOVA sources must appear as rows (main effects cond, par + Total).
    assert any(line.startswith("cond;") for line in lines[1:])
    assert any(line.startswith("par;") for line in lines[1:])
    assert any(line.startswith("Total;") for line in lines[1:])
    # First factor's SS matches the analysis (formatted PT-BR to 6 decimals).
    first = result.anova[0]
    expected_ss = format(first.ss, ".6f").replace(".", ",")
    assert expected_ss in text


def test_csv_coefficients_have_coded_and_natural() -> None:
    result = _golden_result()
    headers, rows = coefficients_table(result)
    assert "Coef. (codificado)" in headers
    assert "Coef. (natural)" in headers
    # Intercept row present.
    assert any(r[0] == "Intercepto" for r in rows)


def test_csv_design_matrix_has_coded_natural_and_responses() -> None:
    factors, responses, runs = _orm_fixture()
    headers, rows = design_matrix_table(factors, responses, runs)
    assert "cond (cod.)" in headers
    assert "par (cod.)" in headers
    assert "DP_50000 (mm)" in headers
    assert len(rows) == len(GOLDEN)
    data = to_csv_ptbr(rows, headers)
    assert data[:3] == b"\xef\xbb\xbf"


# --- XLSX multi-aba ----------------------------------------------------------


def test_xlsx_is_zip_with_three_sheets() -> None:
    result = _golden_result()
    factors, responses, runs = _orm_fixture()
    data = to_xlsx(result, factors, responses, runs)
    assert data[:2] == b"PK"  # zip magic

    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["ANOVA", "Coeficientes", "Matriz de projeto"]


def test_xlsx_values_are_present_and_typed() -> None:
    result = _golden_result()
    factors, responses, runs = _orm_fixture()
    data = to_xlsx(result, factors, responses, runs)
    wb = openpyxl.load_workbook(io.BytesIO(data))

    anova = wb["ANOVA"]
    assert anova["A1"].value == "Fonte"
    sources = [anova.cell(row=r, column=1).value for r in range(2, anova.max_row + 1)]
    assert "cond" in sources
    assert "Total" in sources

    dm = wb["Matriz de projeto"]
    assert dm["A1"].value == "Ensaio"
    # 8 runs -> 9 rows incl. header.
    assert dm.max_row == 1 + len(GOLDEN)
    # A numeric response cell is stored as a number, not a string.
    resp_col = None
    for c in range(1, dm.max_column + 1):
        if dm.cell(row=1, column=c).value == "DP_50000 (mm)":
            resp_col = c
            break
    assert resp_col is not None
    assert isinstance(dm.cell(row=2, column=resp_col).value, int | float)


# --- HTTP end-to-end ---------------------------------------------------------


@pytest_asyncio.fixture
async def client() -> AsyncClient:
    engine = create_async_engine(
        "sqlite+aiosqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    test_session = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async def override_get_session():
        async with test_session() as session:
            try:
                yield session
                await session.commit()
            except Exception:
                await session.rollback()
                raise

    app.dependency_overrides[get_session] = override_get_session
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as async_client:
        yield async_client
    app.dependency_overrides.clear()
    await engine.dispose()


async def _seed_golden(client: AsyncClient) -> int:
    gen = await client.post(
        "/designs/generate",
        json={
            "name": "aula 2^2",
            "replicates": 2,
            "seed": 1,
            "factors": [
                {"name": "cond", "factor_type": "quantitative", "low": 18, "high": 19},
                {"name": "par", "factor_type": "quantitative", "low": 70, "high": 210},
            ],
            "responses": [{"name": "DP_50000", "unit": "mm"}],
        },
    )
    assert gen.status_code == 201
    exp_id = gen.json()["experiment_id"]

    runs = (await client.get(f"/experiments/{exp_id}/runs")).json()["runs"]
    remaining = list(runs)
    values = []
    for cond, par, y in GOLDEN:
        idx = next(
            i
            for i, r in enumerate(remaining)
            if r["factors"]["cond"]["coded"] == cond
            and r["factors"]["par"]["coded"] == par
        )
        run = remaining.pop(idx)
        values.append({"run_id": run["id"], "response_name": "DP_50000", "value": str(y)})
    save = await client.post(f"/experiments/{exp_id}/runs/responses", json={"values": values})
    assert save.status_code == 200
    return exp_id


async def test_export_table_csv_returns_ptbr(client: AsyncClient) -> None:
    exp_id = await _seed_golden(client)
    resp = await client.post(
        "/export/table",
        json={
            "experiment_id": exp_id,
            "response_name": "DP_50000",
            "model": "linear",
            "table": "anova",
        },
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert resp.content[:3] == b"\xef\xbb\xbf"
    assert b";" in resp.content


async def test_export_design_matrix_csv(client: AsyncClient) -> None:
    exp_id = await _seed_golden(client)
    resp = await client.post(
        "/export/table",
        json={
            "experiment_id": exp_id,
            "response_name": "DP_50000",
            "model": "linear",
            "table": "design_matrix",
        },
    )
    assert resp.status_code == 200
    text = resp.content.decode("utf-8-sig")
    assert "cond (cod.)" in text
    assert "DP_50000 (mm)" in text


async def test_export_xlsx_returns_workbook(client: AsyncClient) -> None:
    exp_id = await _seed_golden(client)
    resp = await client.post(
        "/export/xlsx",
        json={
            "experiment_id": exp_id,
            "response_name": "DP_50000",
            "model": "linear",
        },
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert resp.content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["ANOVA", "Coeficientes", "Matriz de projeto"]


async def test_export_table_unknown_table_is_422(client: AsyncClient) -> None:
    exp_id = await _seed_golden(client)
    resp = await client.post(
        "/export/table",
        json={
            "experiment_id": exp_id,
            "response_name": "DP_50000",
            "model": "linear",
            "table": "nope",
        },
    )
    assert resp.status_code == 422


async def test_export_table_unknown_experiment_is_404(client: AsyncClient) -> None:
    resp = await client.post(
        "/export/table",
        json={
            "experiment_id": 999,
            "response_name": "DP_50000",
            "model": "linear",
            "table": "anova",
        },
    )
    assert resp.status_code == 404


# --- PDF report (Playwright) — skipped cleanly if Chromium is missing --------


def _chromium_available() -> bool:
    try:
        from playwright.sync_api import sync_playwright
    except Exception:
        return False
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(args=["--no-sandbox"])
            browser.close()
        return True
    except Exception:
        return False


REPORT_HTML = (
    "<html><body><h1>Relatorio</h1>"
    "<div id='report-ready'>pronto</div></body></html>"
)


@pytest.mark.skipif(
    not _chromium_available(), reason="Chromium do Playwright indisponivel/offline"
)
async def test_build_report_pdf_from_local_html() -> None:
    """Render a minimal file:// page with #report-ready to a real PDF."""
    with tempfile.TemporaryDirectory() as tmp:
        html_path = Path(tmp) / "report.html"
        html_path.write_text(REPORT_HTML, encoding="utf-8")
        url = html_path.as_uri()
        pdf = await build_report_pdf(url, ready_selector="#report-ready")
    assert pdf[:4] == b"%PDF"
    assert len(pdf) > 500
